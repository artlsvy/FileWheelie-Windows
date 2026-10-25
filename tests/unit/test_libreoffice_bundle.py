"""Portable tree/discovery contracts and isolated external-library environment."""

import io
import json
import os
import shutil
import sys
import tempfile
import unittest
from contextlib import contextmanager, redirect_stderr, redirect_stdout
from pathlib import Path
from unittest.mock import Mock, patch

from filewheel import libreoffice_bundle as bundles
from filewheel.converters import ConversionManager, Registry
from filewheel.converters.office import LibreOfficeEngine, run_office


def full_tree(root, *, platform="win32", layout="flat", launcher=None):
    installation = root if layout == "flat" else root / "App/libreoffice"
    program = installation / "program"
    program.mkdir(parents=True)
    files = {
        "soffice.bin": b"binary fixture",
        "types.rdb": b"type fixture",
        "services.rdb": b"service fixture",
    }
    if platform == "win32":
        files.update(
            {
                "soffice.com": b"launcher fixture",
                "soffice.exe": b"launcher fixture",
                "bootstrap.ini": b"bootstrap fixture",
                "fundamental.ini": b"fundamental fixture",
                "mergedlo.dll": b"DLL fixture, not a runnable Windows installation",
            }
        )
    else:
        files.update(
            {
                "soffice": (launcher or "#!/bin/sh\nexit 0\n").encode(),
                "bootstraprc": b"bootstrap fixture",
                "fundamentalrc": b"fundamental fixture",
                "libmergedlo.so": b"library fixture",
            }
        )
    for name, data in files.items():
        (program / name).write_bytes(data)
    if platform != "win32":
        (program / "soffice").chmod(0o755)
    (program / "resource").mkdir()
    (program / "resource/common").write_bytes(b"resource fixture")
    registry = installation / "share/registry"
    registry.mkdir(parents=True)
    (registry / "main.xcd").write_bytes(b"registry fixture")
    (installation / "LICENSE.txt").write_text("LibreOffice license fixture; structural test only")
    return installation


class LibreOfficeBundleTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def test_flat_and_portableapps_complete_windows_tree_prefers_console_launcher(self):
        for layout in ("flat", "portableapps"):
            with self.subTest(layout=layout):
                root = self.root / layout
                installation = full_tree(root, layout=layout)
                location = bundles.validate_bundle(root, platform="win32")
                self.assertEqual(location.source, "bundled")
                self.assertEqual(location.layout, layout)
                self.assertEqual(location.root, installation)
                self.assertEqual(location.executable, installation / "program/soffice.com")
                location.executable.unlink()
                self.assertTrue(
                    bundles.validate_bundle(root, platform="win32").executable.name.endswith(".exe")
                )

    def test_executable_alone_and_missing_required_companions_are_rejected(self):
        program = self.root / "lone/program"
        program.mkdir(parents=True)
        (program / "soffice.exe").write_bytes(b"exe only")
        with self.assertRaises(ValueError):
            bundles.validate_bundle(program.parent, platform="win32")
        for relative in (
            "program/soffice.bin",
            "program/bootstrap.ini",
            "program/types.rdb",
            "program/mergedlo.dll",
            "share/registry/main.xcd",
            "LICENSE.txt",
        ):
            with self.subTest(missing=relative):
                root = self.root / relative.replace("/", "-")
                full_tree(root)
                (root / relative).unlink()
                with self.assertRaises(ValueError):
                    bundles.validate_bundle(root, platform="win32")

    def test_bundle_before_system_and_incomplete_bundle_falls_back(self):
        app = self.root / "application"
        installation = full_tree(app / "tools/libreoffice")
        with patch(
            "filewheel.libreoffice_bundle.shutil.which", return_value="C:/system/soffice.com"
        ) as system:
            selected = bundles.find_libreoffice(app_root=app, platform="win32")
            self.assertEqual(selected.source, "bundled")
            system.assert_not_called()
            (installation / "program/bootstrap.ini").unlink()
            selected = bundles.find_libreoffice(app_root=app, platform="win32")
            self.assertEqual(selected.source, "system")
            self.assertEqual(str(selected.executable), "C:/system/soffice.com")

    def test_incomplete_bundle_cannot_reenter_as_system_path_fallback(self):
        app = self.root / "application"
        program = app / "tools/libreoffice/program"
        program.mkdir(parents=True)
        broken = program / "soffice.com"
        broken.write_bytes(b"incomplete bundled launcher")
        with (
            patch("filewheel.libreoffice_bundle.shutil.which", return_value=str(broken)),
            patch.dict(
                os.environ,
                {
                    "PROGRAMFILES": str(self.root / "missing-programs"),
                    "PROGRAMFILES(X86)": str(self.root / "missing-programs-x86"),
                },
            ),
        ):
            self.assertIsNone(bundles.find_libreoffice(app_root=app, platform="win32"))

    def test_frozen_root_uses_executable_siblings_not_internal(self):
        executable = self.root / "Filewheel.exe"
        executable.write_bytes(b"app fixture")
        installation = full_tree(self.root / "tools/libreoffice")
        with (
            patch.object(sys, "frozen", True, create=True),
            patch.object(sys, "executable", str(executable)),
            patch.object(sys, "_MEIPASS", str(self.root / "_internal"), create=True),
        ):
            selected = bundles.find_libreoffice(platform="win32")
        self.assertEqual(selected.root, installation)

    def test_cli_structural_validation_succeeds_and_reports_invalid_without_traceback(self):
        full_tree(self.root, platform=sys.platform)
        with redirect_stdout(io.StringIO()) as output:
            self.assertEqual(bundles.main(["--validate", str(self.root)]), 0)
        self.assertTrue(json.loads(output.getvalue())["valid"])
        with redirect_stderr(io.StringIO()) as error:
            self.assertEqual(bundles.main(["--validate", str(self.root / "absent")]), 2)
        self.assertNotIn("Traceback", error.getvalue())

    def test_engine_exposes_selected_location_source_and_executable(self):
        app = self.root / "application"
        installation = full_tree(app / "tools/libreoffice", platform=sys.platform)
        with patch("filewheel.libreoffice_bundle.application_root", return_value=app):
            engine = LibreOfficeEngine()
            self.assertEqual(engine.location().source, "bundled")
            self.assertEqual(Path(engine.executable()).parent, installation / "program")

    def test_incomplete_flat_tree_does_not_hide_complete_portableapps_layout(self):
        program = self.root / "program"
        program.mkdir()
        (program / "soffice.exe").write_bytes(b"incomplete launcher")
        full_tree(self.root, layout="portableapps")
        self.assertEqual(
            bundles.validate_bundle(self.root, platform="win32").layout, "portableapps"
        )


class ExternalEnvironmentTests(unittest.TestCase):
    def test_source_environment_is_not_sanitized_or_mutated(self):
        with (
            patch.object(sys, "frozen", False, create=True),
            patch.dict(
                os.environ, {"QT_PLUGIN_PATH": "/custom/plugins", "LD_LIBRARY_PATH": "/custom/lib"}
            ),
        ):
            before = os.environ.copy()
            with bundles.external_tool_environment() as child:
                self.assertEqual(child, before)
                child["ONLY_CHILD"] = "1"
            self.assertEqual(dict(os.environ), before)

    def test_posix_frozen_restores_original_library_path_only_in_child(self):
        with (
            patch.object(sys, "frozen", True, create=True),
            patch.object(sys, "platform", "linux"),
            patch.object(sys, "_MEIPASS", "/app/_internal", create=True),
            patch.dict(
                os.environ,
                {
                    "LD_LIBRARY_PATH": "/app/_internal:/wrong",
                    "LD_LIBRARY_PATH_ORIG": "/system/lib",
                    "QT_PLUGIN_PATH": "/app/_internal/plugins:/user/plugins",
                },
            ),
            bundles.external_tool_environment() as child,
        ):
            self.assertEqual(child["LD_LIBRARY_PATH"], "/system/lib")
            self.assertEqual(child["QT_PLUGIN_PATH"], "/user/plugins")
            self.assertEqual(os.environ["LD_LIBRARY_PATH"], "/app/_internal:/wrong")

    def test_windows_reset_restore_and_qt_path_filter_even_when_spawn_fails(self):
        backend = Mock()
        backend.current_directory.return_value = r"C:\Filewheel\_internal"
        values = {
            "PATH": r"C:\Filewheel\_internal;C:\Windows\System32",
            "QT_PLUGIN_PATH": r"C:\Filewheel\_internal\PySide6;D:\UserPlugins",
            "QT_QPA_PLATFORM_PLUGIN_PATH": r"C:\Filewheel\_internal\platforms",
            "QML2_IMPORT_PATH": r"C:\Filewheel\_internal\qml",
        }
        with (
            patch.object(sys, "frozen", True, create=True),
            patch.object(sys, "platform", "win32"),
            patch.object(sys, "_MEIPASS", r"C:\Filewheel\_internal", create=True),
            patch.dict(os.environ, values),
        ):
            with (
                self.assertRaisesRegex(OSError, "spawn failed"),
                bundles.external_tool_environment(_windows_backend=backend) as child,
            ):
                backend.set_directory.assert_called_once_with(None)
                self.assertEqual(child["PATH"], r"C:\Windows\System32")
                self.assertEqual(child["QT_PLUGIN_PATH"], r"D:\UserPlugins")
                self.assertNotIn("QT_QPA_PLATFORM_PLUGIN_PATH", child)
                self.assertNotIn("QML2_IMPORT_PATH", child)
                raise OSError("spawn failed")
            self.assertEqual(os.environ["QT_PLUGIN_PATH"], values["QT_PLUGIN_PATH"])
        self.assertEqual(
            [call.args[0] for call in backend.set_directory.call_args_list],
            [None, r"C:\Filewheel\_internal"],
        )

    def test_process_spawn_scope_ends_before_communicate(self):
        state = {"inside": False}

        @contextmanager
        def context():
            state["inside"] = True
            try:
                yield {"SAFE_CHILD": "1"}
            finally:
                state["inside"] = False

        def communicate(*, timeout):
            self.assertFalse(state["inside"])
            return b"", b""

        process = Mock(returncode=0)
        process.communicate.side_effect = communicate

        def spawn(*args, **kwargs):
            self.assertTrue(state["inside"])
            self.assertEqual(kwargs["env"], {"SAFE_CHILD": "1"})
            self.assertFalse(kwargs["shell"])
            return process

        with (
            patch("filewheel.converters.office.external_tool_environment", context),
            patch("filewheel.converters.office.subprocess.Popen", side_effect=spawn),
        ):
            self.assertEqual(run_office(["soffice"], workdir=Path.cwd(), timeout=1), 0)

    def test_restore_failure_after_successful_spawn_stops_owned_child(self):
        @contextmanager
        def context():
            yield {"SAFE_CHILD": "1"}
            raise OSError("restore failed")

        process = Mock()
        with (
            patch("filewheel.converters.office.external_tool_environment", context),
            patch("filewheel.converters.office.subprocess.Popen", return_value=process),
            patch("filewheel.converters.office._stop_office_tree") as stop,
            self.assertRaisesRegex(OSError, "restore failed"),
        ):
            run_office(["soffice"], workdir=Path.cwd(), timeout=1)
        stop.assert_called_once_with(process, {"SAFE_CHILD": "1"})

    def test_windows_dll_reset_scopes_are_serialized(self):
        import time
        from concurrent.futures import ThreadPoolExecutor

        backend = Mock()
        backend.current_directory.return_value = r"C:\Filewheel\_internal"
        entries = []

        def spawn_scope(index):
            with bundles.external_tool_environment(_windows_backend=backend):
                entries.append((index, "enter"))
                time.sleep(0.01)
                entries.append((index, "exit"))

        with (
            patch.object(sys, "frozen", True, create=True),
            patch.object(sys, "platform", "win32"),
            ThreadPoolExecutor(max_workers=2) as pool,
        ):
            list(pool.map(spawn_scope, (1, 2)))
        self.assertEqual([phase for _, phase in entries], ["enter", "exit", "enter", "exit"])
        self.assertEqual(
            [call.args[0] for call in backend.set_directory.call_args_list],
            [None, r"C:\Filewheel\_internal", None, r"C:\Filewheel\_internal"],
        )


class BundledOfficeIntegrationTests(unittest.TestCase):
    def test_actual_linux_wrapper_selected_as_bundle_renders_valid_pdf(self):
        if sys.platform != "linux" or shutil.which("libreoffice") is None:
            self.skipTest("Installed Linux LibreOffice required for wrapper integration")
        try:
            import pymupdf
            from docx import Document
        except ImportError:
            self.skipTest("PyMuPDF/python-docx fixture dependencies unavailable")
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            app = root / "application"
            executable = shutil.which("libreoffice")
            # This fixture verifies discovery and process integration by
            # delegating to installed LO; it is not a self-contained binary tree.
            full_tree(
                app / "tools/libreoffice",
                platform="linux",
                launcher=f'#!/bin/sh\nexec "{executable}" "$@"\n',
            )
            source = root / "Unicode æ fixture.docx"
            document = Document()
            document.add_heading("Portable selection fixture", 1)
            document.add_paragraph("Actual editable text rendered through selected local engine")
            document.save(source)
            original = source.read_bytes()
            with patch("filewheel.libreoffice_bundle.application_root", return_value=app):
                engine = LibreOfficeEngine(timeout=30)
                self.assertEqual(engine.location().source, "bundled")
                result = ConversionManager(Registry((engine,))).convert(source, ".pdf")
            with pymupdf.open(result.path) as pdf:
                self.assertIn("Actual editable text", "".join(page.get_text() for page in pdf))
            self.assertEqual(source.read_bytes(), original)
            self.assertFalse(list(root.glob(".filewheel-*")))
