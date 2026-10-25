"""Actual local pdf2docx runs on generated text, table, scanned and corrupt PDFs."""

import hashlib
import importlib.util
import io
import tempfile
import unittest
from pathlib import Path
from threading import Event
from unittest.mock import patch
from zipfile import ZipFile

from filewheel.converters import (
    ConversionCancelled,
    ConversionError,
    ConversionManager,
    EngineOutput,
    Registry,
)
from filewheel.converters.pdf_docx import PdfDocxEngine

AVAILABLE = all(importlib.util.find_spec(name) is not None for name in ("pdf2docx", "pymupdf", "docx"))


@unittest.skipUnless(AVAILABLE, "Optional local PDF→DOCX dependencies unavailable")
class PdfDocxTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.folder = Path(self.temporary.name)
        self.engine = PdfDocxEngine()
        self.manager = ConversionManager(Registry((self.engine,)))

    def tearDown(self):
        self.temporary.cleanup()

    def make_pdf(self, *, image_page=False, mixed=False, table=False):
        import pymupdf
        source = self.folder / "Editable résumé.pdf"
        with pymupdf.open() as document:
            if not image_page or mixed:
                page = document.new_page()
                page.insert_text((50, 55), "Quarterly Report", fontsize=18)
                page.insert_text((50, 90), "Editable document text survives conversion.", fontsize=11)
                page.insert_text((50, 115), "First bullet item", fontsize=11)
                page.insert_text((50, 135), "Second bullet item", fontsize=11)
                if table:
                    for y in (170, 200, 230):
                        page.draw_line((50, y), (350, y))
                    for x in (50, 200, 350):
                        page.draw_line((x, 170), (x, 230))
                    for point, text in (((60, 190), "Item"), ((210, 190), "Value"),
                                        ((60, 220), "Widgets"), ((210, 220), "42")):
                        page.insert_text(point, text, fontsize=11)
            if image_page:
                from PIL import Image
                image = io.BytesIO()
                with Image.new("RGB", (40, 40), "white") as pixels:
                    pixels.save(image, format="PNG")
                page = document.new_page()
                page.insert_image(pymupdf.Rect(50, 50, 250, 250), stream=image.getvalue())
            document.save(source)
        return source

    def text(self, path):
        from docx import Document
        document = Document(path)
        return ("\n".join(paragraph.text for paragraph in document.paragraphs)
                + "\n" + "\n".join(cell.text for table in document.tables
                                      for row in table.rows for cell in row.cells))

    def test_editable_text_and_basic_table_with_source_preserved(self):
        source = self.make_pdf(table=True)
        digest = hashlib.sha256(source.read_bytes()).digest()
        result = self.manager.convert(source, ".docx")
        text = self.text(result.path)
        for phrase in ("Quarterly Report", "Editable document text", "First bullet item", "Widgets", "42"):
            self.assertIn(phrase, text)
        from docx import Document
        self.assertGreaterEqual(len(Document(result.path).tables), 1)
        self.assertTrue(result.warnings)
        self.assertEqual(hashlib.sha256(source.read_bytes()).digest(), digest)
        self.assertFalse(list(self.folder.glob(".filewheel-*")))

    def test_scanned_pdf_rejects_with_ocr_explanation(self):
        source = self.make_pdf(image_page=True)
        with self.assertRaisesRegex(ConversionError, "OCR"):
            self.manager.convert(source, ".docx")
        self.assertEqual(list(self.folder.iterdir()), [source])

    def test_mixed_pdf_warns_about_noneditable_pages(self):
        source = self.make_pdf(image_page=True, mixed=True)
        result = self.manager.convert(source, ".docx")
        self.assertIn("Quarterly Report", self.text(result.path))
        self.assertTrue(any("no extractable text" in warning for warning in result.warnings))

    def test_corrupt_pdf_and_encrypted_pdf_reject(self):
        import pymupdf
        source = self.folder / "corrupt.pdf"
        source.write_bytes(b"invalid pdf")
        with self.assertRaisesRegex(ConversionError, "corrupt"):
            self.manager.convert(source, ".docx")
        with pymupdf.open() as document:
            document.new_page().insert_text((50, 55), "Secret")
            document.save(source, encryption=pymupdf.PDF_ENCRYPT_AES_256,
                          owner_pw="owner", user_pw="password")
        with self.assertRaisesRegex(ConversionError, "Password"):
            self.manager.convert(source, ".docx")
        self.assertFalse(list(self.folder.glob("*.docx")))

    def test_input_limits_and_cancellation_publish_nothing(self):
        source = self.make_pdf(image_page=True, mixed=True)
        for engine in (PdfDocxEngine(max_pages=1), PdfDocxEngine(max_bytes=1)):
            with self.assertRaisesRegex(ConversionError, "limit"):
                ConversionManager(Registry((engine,))).convert(source, ".docx")
        event = Event()
        def progress(value, message):
            if value >= .1:
                event.set()
        with self.assertRaises(ConversionCancelled):
            self.manager.convert(source, ".docx", progress=progress, cancel=event)
        self.assertEqual(list(self.folder.iterdir()), [source])

    def test_docx_validation_rejects_archive_without_editable_text(self):
        path = self.folder / "output.docx"
        with ZipFile(path, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types />")
            archive.writestr("word/document.xml", "<document />")
        with self.assertRaisesRegex(ConversionError, "editable text"):
            self.engine.validate(EngineOutput(path, "output.docx"))

    def test_unavailable_backend_not_registered(self):
        with patch("filewheel.converters.pdf_docx.find_spec", return_value=None):
            self.assertEqual(Registry((self.engine,)).capabilities, ())

    def test_failed_page_conversion_closes_backend_and_cleans_partial(self):
        source = self.make_pdf()
        with patch("pdf2docx.Converter") as converter:
            converter.return_value.convert.side_effect = RuntimeError("page parse")
            with self.assertRaisesRegex(ConversionError, "reconstruction"):
                self.manager.convert(source, ".docx")
            converter.return_value.close.assert_called_once()
        self.assertEqual(list(self.folder.iterdir()), [source])

    def test_cancel_during_backend_call_is_observed_before_publication(self):
        source = self.make_pdf()
        event = Event()
        with patch("pdf2docx.Converter") as converter:
            converter.return_value.convert.side_effect = lambda *args, **kwargs: event.set()
            with self.assertRaises(ConversionCancelled):
                self.manager.convert(source, ".docx", cancel=event)
            converter.return_value.close.assert_called_once()
        self.assertEqual(list(self.folder.iterdir()), [source])
