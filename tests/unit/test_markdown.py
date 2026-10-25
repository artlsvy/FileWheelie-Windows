"""Real local MarkItDown extraction from generated supported-format fixtures."""

import hashlib
import importlib.util
import io
import os
import sys
import tempfile
import unittest
from pathlib import Path
from threading import Event
from unittest.mock import patch

# Prevent ONNX telemetry before optional dependency imports in this test module.
os.environ["ORT_DISABLE_TELEMETRY"] = "1"

from filewheel.converters import ConversionCancelled, ConversionError, ConversionManager, Registry
from filewheel.converters.markdown import MarkItDownEngine, RejectRequestsSession

AVAILABLE = sys.version_info < (3, 15) and all(
    importlib.util.find_spec(module) is not None
    for module in ("markitdown", "docx", "pptx", "openpyxl", "pymupdf", "mammoth", "pandas")
)


@unittest.skipUnless(
    AVAILABLE, "Optional MarkItDown fixtures require Python<3.15 and document dependencies"
)
class MarkdownTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.folder = Path(self.temporary.name)
        self.engine = MarkItDownEngine()
        self.manager = ConversionManager(Registry((self.engine,)))

    def tearDown(self):
        self.temporary.cleanup()

    def convert(self, source):
        original = hashlib.sha256(source.read_bytes()).digest()
        result = self.manager.convert(source, ".md")
        text = result.path.read_text(encoding="utf-8")
        self.assertEqual(hashlib.sha256(source.read_bytes()).digest(), original)
        self.assertTrue(result.warnings)
        self.assertFalse(list(self.folder.glob(".filewheel-*")))
        return text

    def test_docx_heading_lists_table_and_unicode(self):
        from docx import Document

        source = self.folder / "Résumé document.docx"
        document = Document()
        document.add_heading("Résumé heading", 1)
        document.add_paragraph("Useful list item", style="List Bullet")
        table = document.add_table(rows=2, cols=2)
        for cell, text in zip(
            (cell for row in table.rows for cell in row.cells), ("Item", "Value", "Widgets", "42")
        ):
            cell.text = text
        document.save(source)
        text = self.convert(source)
        self.assertIn("# Résumé heading", text)
        self.assertIn("Useful list item", text)
        self.assertIn("Widgets", text)
        self.assertIn("|", text)

    def test_pptx_title_and_body(self):
        from pptx import Presentation

        source = self.folder / "slides.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Résumé slides"
        slide.placeholders[1].text = "Useful slide content"
        presentation.save(source)
        text = self.convert(source)
        self.assertIn("Résumé slides", text)
        self.assertIn("Useful slide content", text)

    def test_xlsx_sheet_table_and_unicode(self):
        from openpyxl import Workbook

        source = self.folder / "sheet.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Résumé data"
        sheet.append(["Name", "Count"])
        sheet.append(["Widgets", 42])
        workbook.save(source)
        workbook.close()
        text = self.convert(source)
        self.assertIn("Résumé data", text)
        self.assertIn("Widgets", text)
        self.assertIn("42", text)
        self.assertIn("|", text)

    def test_html_heading_list_and_remote_images_do_not_fetch(self):
        source = self.folder / "page.html"
        source.write_text(
            "<!doctype html><html><body><h1>Résumé heading</h1>"
            "<ul><li>Useful list item</li></ul>"
            '<img src="https://invalid.example/private.png">'
            '<a href="https://invalid.example/">Link</a></body></html>',
            encoding="utf-8",
        )
        with patch("socket.create_connection", side_effect=AssertionError("network attempted")):
            text = self.convert(source)
        self.assertIn("# Résumé heading", text)
        self.assertIn("Useful list item", text)
        self.assertIn("https://invalid.example/", text)

    def test_pdf_extractable_text_and_scanned_pdf_rejection(self):
        import pymupdf
        from PIL import Image

        source = self.folder / "document.pdf"
        with pymupdf.open() as document:
            document.new_page().insert_text((50, 60), "Useful PDF document text")
            document.save(source)
        self.assertIn("Useful PDF document text", self.convert(source))
        scanned = self.folder / "scanned.pdf"
        image = io.BytesIO()
        with Image.new("RGB", (40, 40), "white") as pixels:
            pixels.save(image, format="PNG")
        with pymupdf.open() as document:
            document.new_page().insert_image(pymupdf.Rect(0, 0, 100, 100), stream=image.getvalue())
            document.save(scanned)
        with self.assertRaisesRegex(ConversionError, "OCR"):
            self.manager.convert(scanned, ".md")
        self.assertFalse(scanned.with_suffix(".md").exists())

    def test_corrupt_or_wrong_format_documents_leave_no_outputs(self):
        for suffix in (".docx", ".pptx", ".xlsx", ".pdf", ".html"):
            with self.subTest(suffix=suffix):
                source = self.folder / ("corrupt" + suffix)
                source.write_bytes(b"not a document")
                with self.assertRaises(ConversionError):
                    self.manager.convert(source, ".md")
                self.assertFalse(source.with_suffix(".md").exists())
                self.assertFalse(list(self.folder.glob(".filewheel-*")))

    def test_cancellation_before_extraction_and_request_session_denies_network(self):
        source = self.folder / "page.html"
        source.write_text("<h1>Heading</h1>", encoding="utf-8")
        event = Event()
        event.set()
        with self.assertRaises(ConversionCancelled):
            self.manager.convert(source, ".md", cancel=event)
        with self.assertRaisesRegex(ConversionError, "Network"):
            RejectRequestsSession().get("https://invalid.example/")

    def test_python315_and_missing_format_dependencies_hide_capabilities(self):
        with patch("filewheel.converters.markdown.sys.version_info", (3, 15, 0)):
            self.assertEqual(MarkItDownEngine().capabilities, ())
        real_find = importlib.util.find_spec
        with patch(
            "filewheel.converters.markdown.find_spec",
            side_effect=lambda name: None if name == "mammoth" else real_find(name),
        ):
            extensions = {c.input_extension for c in MarkItDownEngine().capabilities}
        self.assertNotIn(".docx", extensions)
        self.assertIn(".html", extensions)

    def test_missing_pdfplumber_hides_only_pdf_pair(self):
        real_find = importlib.util.find_spec
        with patch("filewheel.converters.markdown.find_spec",
                   side_effect=lambda name: None if name == "pdfplumber" else real_find(name)):
            extensions = {c.input_extension for c in MarkItDownEngine().capabilities}
        self.assertNotIn(".pdf", extensions)
        self.assertIn(".html", extensions)
