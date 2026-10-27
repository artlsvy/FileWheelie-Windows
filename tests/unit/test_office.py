import subprocess
import tempfile
import unittest
from pathlib import Path
from threading import Event
from unittest.mock import patch

from filewheel.converters import ConversionCancelled, ConversionError, ConversionManager, Registry
from filewheel.converters.office import LibreOfficeEngine


class OfficeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.source = self.folder / 'Unicode æ notes.docx'
        try:
            from docx import Document
        except ImportError:
            self.skipTest('python-docx unavailable')
        document = Document()
        document.add_heading('Office fixture', 0)
        document.add_paragraph('Readable local document text')
        document.add_table(rows=1, cols=2).cell(0, 0).text = 'Cell content'
        document.save(self.source)
        self.engine = LibreOfficeEngine(timeout=30)

    def test_actual_libreoffice_output_and_collision(self):
        if not self.engine.available():
            self.skipTest('LibreOffice/PyMuPDF unavailable')
        import pymupdf
        manager = ConversionManager(Registry((self.engine,)))
        before = self.source.read_bytes()
        existing = self.source.with_suffix('.pdf')
        existing.write_bytes(b'existing output')
        result = manager.convert(self.source, '.pdf')
        self.assertEqual(result.path.name, 'Unicode æ notes (2).pdf')
        with pymupdf.open(result.path) as output:
            self.assertIn('Readable local document text', ''.join(page.get_text() for page in output))
        self.assertEqual(existing.read_bytes(), b'existing output')
        self.assertEqual(self.source.read_bytes(), before)
        self.assertFalse(list(self.folder.glob('.filewheel-*')))

    def test_timeout_cleanup_and_safe_arguments(self):
        with (patch.object(self.engine, 'executable', return_value='/fake/soffice'),
              patch('filewheel.converters.office.find_spec', return_value=object()),
              patch('filewheel.converters.office.run_office',
                    side_effect=ConversionError('LibreOffice timed out')) as run):
            manager = ConversionManager(Registry((self.engine,)))
            with self.assertRaisesRegex(ConversionError, 'timed out'):
                manager.convert(self.source, '.pdf')
        self.assertEqual(run.call_args.kwargs['timeout'], 30)
        self.assertIn('cancel', run.call_args.kwargs)
        self.assertTrue(any('UserInstallation=file:' in arg for arg in run.call_args.args[0]))
        self.assertEqual(list(self.folder.iterdir()), [self.source])

    def test_missing_engine_is_hidden(self):
        with patch.object(self.engine, 'executable', return_value=None):
            self.assertEqual(Registry((self.engine,)).capabilities, ())

    def test_corrupt_docx_never_invokes_process(self):
        self.source.write_bytes(b'corrupt')
        with (patch.object(self.engine, 'available', return_value=True),
              patch.object(self.engine, 'executable', return_value='/fake/soffice'),
              patch('filewheel.converters.office.run_office') as run,
              self.assertRaises(ConversionError)):
            ConversionManager(Registry((self.engine,))).convert(self.source, '.pdf')
        run.assert_not_called()
        self.assertEqual(list(self.folder.iterdir()), [self.source])

    def test_cancel_stops_owned_office_tree_between_bounded_waits(self):
        from unittest.mock import Mock

        from filewheel.converters.office import run_office
        event = Event()
        process = Mock(pid=123)
        def waiting(*, timeout):
            self.assertLessEqual(timeout, .1)
            event.set()
            raise subprocess.TimeoutExpired('soffice', timeout)
        process.communicate.side_effect = waiting
        with (patch('filewheel.converters.office.subprocess.Popen', return_value=process),
              patch('filewheel.converters.office._stop_office_tree') as stop,
              self.assertRaises(ConversionCancelled)):
            run_office(['soffice'], workdir=self.folder, timeout=30, cancel=event)
        stop.assert_called_once()
        self.assertIs(stop.call_args.args[0], process)
        self.assertEqual(process.communicate.call_count, 1)

    def test_precancelled_office_never_starts_process(self):
        from filewheel.converters.office import run_office
        event = Event()
        event.set()
        with (patch('filewheel.converters.office.subprocess.Popen') as spawn,
              self.assertRaises(ConversionCancelled)):
            run_office(['soffice'], workdir=self.folder, timeout=30, cancel=event)
        spawn.assert_not_called()

    def test_real_wrapper_timeout_stops_descendant(self):
        import sys
        import time

        from filewheel.converters.office import run_office
        script = self.folder / 'wrapper.py'
        pidfile = self.folder / 'child.pid'
        script.write_text("import subprocess,sys,time\nfrom pathlib import Path\n"
                          "p=subprocess.Popen([sys.executable,'-c','import time; time.sleep(60)'])\n"
                          "Path(sys.argv[1]).write_text(str(p.pid))\ntime.sleep(60)\n")
        started = time.monotonic()
        with self.assertRaisesRegex(ConversionError, 'timed out'):
            run_office([sys.executable, str(script), str(pidfile)],
                       workdir=self.folder, timeout=1)
        self.assertLess(time.monotonic() - started, 8)
        child = int(pidfile.read_text())
        import ctypes
        from ctypes import wintypes
        kernel = ctypes.WinDLL('kernel32', use_last_error=True)
        kernel.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
        kernel.OpenProcess.restype = wintypes.HANDLE
        kernel.WaitForSingleObject.argtypes = [wintypes.HANDLE, wintypes.DWORD]
        kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        def running():
            handle = kernel.OpenProcess(0x00100000, False, child)  # SYNCHRONIZE
            if not handle:
                if ctypes.get_last_error() == 87:  # Process no longer exists.
                    return False
                raise ctypes.WinError(ctypes.get_last_error())
            try:
                return kernel.WaitForSingleObject(handle, 0) == 0x102
            finally:
                kernel.CloseHandle(handle)
        deadline = time.monotonic() + 2
        while running() and time.monotonic() < deadline:
            time.sleep(.01)
        self.assertFalse(running())

    def test_real_wrapper_cancellation_finishes_promptly(self):
        import sys
        import time
        from threading import Timer

        from filewheel.converters.office import run_office
        event = Event()
        timer = Timer(.2, event.set)
        timer.start()
        self.addCleanup(timer.cancel)
        started = time.monotonic()
        with self.assertRaises(ConversionCancelled):
            run_office([sys.executable, '-c', 'import time; time.sleep(60)'],
                       workdir=self.folder, timeout=30, cancel=event)
        self.assertLess(time.monotonic() - started, 3)
