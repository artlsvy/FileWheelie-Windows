"""Real raster-PDF output plus isolated dependency, cancellation and platform guards."""

import hashlib
import shutil
import sys
from pathlib import Path
from threading import Event, Timer
from urllib.parse import unquote

import pymupdf
import pytest
from docx import Document

from filewheel.converters import ConversionManager, local_ocr, tesseract_ocr
from filewheel.converters.base import ConversionCancelled, ConversionError, DependencyUnavailable
from filewheel.converters.office import OfficeExit


def raster_pdf(path, *, text=True, pages=1):
    with pymupdf.open() as raster, pymupdf.open() as document:
        original = raster.new_page(width=600, height=180)
        if text:
            original.insert_text((25, 70), "Filewheel local OCR example text", fontsize=24)
        else:
            original.draw_circle((300, 90), 45, color=(0.1, 0.5, 0.2), fill=(0.1, 0.5, 0.2))
        pixels = original.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes("png")
        for _ in range(pages):
            page = document.new_page(width=600, height=180)
            page.insert_image(page.rect, stream=pixels)
        document.save(path)
    with pymupdf.open(path) as document:
        assert not any(page.get_text().strip() for page in document)
    return path


@pytest.mark.parametrize("extension", [".md", ".docx"])
def test_real_scanned_pdf_unicode_collisions_and_output(tmp_path, extension):
    assert shutil.which("tesseract"), "Install Tesseract for the real Linux OCR tests"
    source = raster_pdf(tmp_path / "Scanned æøå file.pdf", pages=2)
    before = hashlib.sha256(source.read_bytes()).digest()
    manager = ConversionManager()
    result = manager.convert(source, extension)
    if extension == ".docx":
        text = "\n".join(p.text for p in Document(result.path).paragraphs)
        assert len(Document(result.path).inline_shapes) == 2
    else:
        text = result.path.read_text(encoding="utf-8")
        assert len(result.outputs) == 3
        assert "## Page 2" in text
        for image in result.outputs[1:]:
            assert image.exists() and image.name in unquote(text)
    assert "Filewheel" in text and "example text" in text
    assert any("Tesseract" in warning for warning in result.warnings)
    first_output = result.path.read_bytes()
    again = manager.convert(source, extension)
    assert again.path != result.path and result.path.read_bytes() == first_output
    assert hashlib.sha256(source.read_bytes()).digest() == before
    assert not list(tmp_path.glob(".filewheel-*"))


def test_picture_without_letters_preserves_markdown_images(tmp_path):
    source = raster_pdf(tmp_path / "photos.pdf", text=False)
    result = ConversionManager().convert(source, ".md")
    assert len(result.outputs) == 2
    assert result.outputs[1].name in result.path.read_text()
    assert any("No text was recognized" in warning for warning in result.warnings)


def test_text_pdf_does_not_need_ocr(tmp_path, monkeypatch):
    def forbidden(*args):
        raise AssertionError("OCR must not run for a text-layer PDF")

    monkeypatch.setattr(tesseract_ocr, "location", forbidden)
    source = tmp_path / "normal.pdf"
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((50, 60), "Useful text without OCR")
        pdf.save(source)
    result = ConversionManager().convert(source, ".md")
    assert "Useful text" in result.path.read_text()


def test_bundled_discovery_data_and_missing_dependency(tmp_path, monkeypatch):
    monkeypatch.setattr(tesseract_ocr, "bundled_tool_root", lambda: tmp_path)
    monkeypatch.setenv("FILEWHEEL_BUNDLED_TOOLS_ONLY", "1")
    monkeypatch.setattr(tesseract_ocr.shutil, "which", lambda _: "/host/tesseract")
    with pytest.raises(DependencyUnavailable, match="local OCR"):
        tesseract_ocr.location()
    tool = tmp_path / "tools/tesseract/tesseract"
    tool.parent.mkdir(parents=True)
    tool.write_text("#!/bin/sh\nexit 0\n")
    tool.chmod(0o755)
    with pytest.raises(DependencyUnavailable, match="recognition data"):
        tesseract_ocr.location()
    data = tool.parent / "tessdata"
    data.mkdir()
    (data / "eng.traineddata").write_text("fixture")
    assert tesseract_ocr.location() == (tool, data, "bundled")


def test_failed_ocr_leaves_no_outputs_or_source_changes(tmp_path, monkeypatch):
    source = raster_pdf(tmp_path / "scan.pdf")
    before = source.read_bytes()
    monkeypatch.setattr(
        tesseract_ocr,
        "location",
        lambda: (_ for _ in ()).throw(DependencyUnavailable("OCR missing")),
    )
    with pytest.raises(DependencyUnavailable):
        ConversionManager().convert(source, ".md")
    assert source.read_bytes() == before and set(tmp_path.iterdir()) == {source}


@pytest.mark.parametrize("response", [b"\xff", b"binary\0text", b"x" * (2 * 1024 * 1024 + 1), None])
def test_invalid_ocr_output_rejected_and_cleaned(tmp_path, monkeypatch, response):
    monkeypatch.setattr(tesseract_ocr, "location", lambda: (Path("/bin/tesseract"), None, "system"))

    def run(arguments, **kwargs):
        assert kwargs["timeout"] == 60 and kwargs["tool_name"] == "Tesseract OCR"
        if response is not None:
            (tmp_path / "ocr-result.txt").write_bytes(response)
        return OfficeExit(0)

    monkeypatch.setattr(tesseract_ocr, "run_office", run)
    with pytest.raises(ConversionError, match="invalid result"):
        tesseract_ocr.recognize_image(tmp_path / "page.png", tmp_path, None)
    assert not (tmp_path / "ocr-result.txt").exists()


def test_failed_process_and_cancellation(tmp_path, monkeypatch):
    tool = tmp_path / "slow-ocr"
    tool.write_text("#!/usr/bin/python3\nimport time\ntime.sleep(30)\n")
    tool.chmod(0o755)
    monkeypatch.setattr(tesseract_ocr, "location", lambda: (tool, None, "system"))
    cancel = Event()
    timer = Timer(0.2, cancel.set)
    timer.start()
    try:
        with pytest.raises(ConversionCancelled):
            tesseract_ocr.recognize_image(tmp_path / "page.png", tmp_path, cancel)
    finally:
        timer.cancel()
    monkeypatch.setattr(
        tesseract_ocr, "run_office", lambda *a, **kw: OfficeExit(1, stderr=b"test failure")
    )
    with pytest.raises(ConversionError, match="test failure"):
        tesseract_ocr.recognize_image(tmp_path / "page.png", tmp_path, None)


def test_language_dependency_and_timeout_are_clear(tmp_path, monkeypatch):
    monkeypatch.setattr(tesseract_ocr, "location", lambda: (Path("/bin/tesseract"), None, "system"))
    monkeypatch.setattr(
        tesseract_ocr,
        "run_office",
        lambda *a, **kw: OfficeExit(1, stderr=b"Failed loading language eng"),
    )
    with pytest.raises(DependencyUnavailable, match="recognition data"):
        tesseract_ocr.recognize_image(tmp_path / "page.png", tmp_path, None)

    def timeout(*args, **kwargs):
        raise ConversionError("Tesseract OCR timed out; conversion stopped")

    monkeypatch.setattr(tesseract_ocr, "run_office", timeout)
    with pytest.raises(ConversionError, match="timed out"):
        tesseract_ocr.recognize_image(tmp_path / "page.png", tmp_path, None)


def test_windows_recognition_still_uses_native_helper(tmp_path, monkeypatch):
    from filewheel.converters import windows_ocr

    monkeypatch.setattr(windows_ocr, "recognize_image", lambda *args: (["native text"], "en-US"))
    monkeypatch.setattr(sys, "platform", "win32")
    assert local_ocr.recognize_image(tmp_path / "a.png", tmp_path, None) == (
        ["native text"],
        "en-US",
    )
    assert local_ocr.backend_name() == "Windows"
