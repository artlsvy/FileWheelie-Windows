"""Measure local generated fixtures: direct engine transaction vs fresh worker.

Not a converter speed ranking or user-document benchmark. Uses existing engines
and their validation/publication behavior without changing application code.
"""
import argparse
import hashlib
import json
import statistics
import sys
import tempfile
from pathlib import Path
from threading import Event
from time import monotonic


def measure(source, extension, engine, samples):
    from filewheel.converters import ConversionManager, Registry
    from filewheel.converters.process_worker import run_in_process
    manager = ConversionManager(Registry((engine,)))
    if not manager.registry.capabilities:
        return {'engine': engine.engine_id, 'status': 'dependency_unavailable'}
    original = hashlib.sha256(source.read_bytes()).hexdigest()
    direct, isolated = [], []
    for values, worker in ((direct, False), (isolated, True)):
        for _ in range(samples):
            start = monotonic()
            if worker:
                result = run_in_process(manager, source, extension, progress=lambda *args: None,
                                        cancel=Event())
            else:
                result = manager.convert(source, extension)
            values.append((monotonic() - start) * 1000)
            if not result.outputs or not all(path.stat().st_size for path in result.outputs):
                raise RuntimeError('Benchmark did not produce validated output')
    if hashlib.sha256(source.read_bytes()).hexdigest() != original:
        raise RuntimeError('Fixture input changed')
    if list(source.parent.glob('.filewheel-*')):
        raise RuntimeError('Fixture staging was not cleaned')
    return {'engine': engine.engine_id, 'input_extension': source.suffix, 'output_extension': extension,
            'samples': samples, 'status': 'ok', 'source_unchanged': True,
            'direct_transaction_ms': direct, 'isolated_transaction_ms': isolated,
            'direct_median_ms': statistics.median(direct),
            'isolated_median_ms': statistics.median(isolated)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--samples', type=int, default=3)
    args = parser.parse_args()
    if not 1 <= args.samples <= 10:
        parser.error('samples must be between1 and10')
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
    import pymupdf
    from docx import Document
    from PIL import Image

    from filewheel.converters.images import PillowEngine
    from filewheel.converters.markdown import MarkItDownEngine
    from filewheel.converters.pdf_images import PdfImagesEngine
    with tempfile.TemporaryDirectory(prefix='filewheel-conversion-benchmark-') as directory:
        root = Path(directory)
        png = root / 'image.png'
        Image.new('RGBA', (1920, 1080), (30, 80, 130, 140)).save(png)
        pdf = root / 'three-pages.pdf'
        with pymupdf.open() as document:
            for index in range(3):
                document.new_page().insert_text((50, 60), f'Local generated benchmark page {index + 1}')
            document.save(pdf)
        docx = root / 'simple.docx'
        document = Document()
        document.add_heading('Generated benchmark', 1)
        document.add_paragraph('A short local document; no external resources.')
        document.save(docx)
        results = [measure(png, '.jpg', PillowEngine(), args.samples),
                   measure(pdf, '.png', PdfImagesEngine(), args.samples),
                   measure(docx, '.md', MarkItDownEngine(), args.samples)]
    payload = {'python': sys.version.split()[0], 'platform': sys.platform, 'results': results,
               'limits': ['Generated simple fixtures only; libraries used to create them are imported in parent',
                          'Direct measures encode/validate/publish; isolated also includes fresh interpreter/import overhead',
                          'Production GUI always uses isolated workers, never this direct benchmark path',
                          'No Windows or complex/scanned PDF performance is established']}
    with args.output.open('x', encoding='utf-8') as stream:
        json.dump(payload, stream, indent=2)
    for row in results:
        print(row['engine'], row.get('direct_median_ms'), row.get('isolated_median_ms'), row['status'])


if __name__ == '__main__':
    main()
