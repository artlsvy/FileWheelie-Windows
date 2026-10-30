"""Shared bounded PDF OCR and editable/reference output; recognition stays offline."""

import sys

from filewheel.converters.base import ConversionError, EngineOutput, check_cancel


def backend_name():
    return "Windows" if sys.platform == "win32" else "Tesseract"


def recognize_image(image, workdir, cancel):
    if sys.platform == "win32":
        from filewheel.converters.windows_ocr import recognize_image as recognize
    else:
        from filewheel.converters.tesseract_ocr import recognize_image as recognize
    return recognize(image, workdir, cancel)


def ocr_pdf_pages(source, workdir, progress, cancel):
    """Yield bounded page previews and locally recognized lines."""
    import pymupdf

    total_bytes = 0
    with pymupdf.open(source) as pdf:
        for index, page in enumerate(pdf):
            check_cancel(cancel)
            progress(
                0.25 + 0.5 * index / len(pdf),
                f"Reading page {index + 1} of {len(pdf)} with local OCR",
            )
            scale = min(3, 3000 / max(page.rect.width, page.rect.height))
            image = workdir / f"ocr-page-{index + 1}.png"
            page.get_pixmap(
                matrix=pymupdf.Matrix(scale, scale), colorspace=pymupdf.csRGB, alpha=False
            ).save(image)
            total_bytes += image.stat().st_size
            if total_bytes > 256 * 1024 * 1024:
                raise ConversionError("PDF exceeds the local OCR image size limit")
            lines, language = recognize_image(image, workdir, cancel)
            yield index + 1, lines, language, image


def image_pdf_to_docx(source, workdir, progress, cancel):
    """Simple editable paragraphs plus source-page references, not layout reconstruction."""
    import pymupdf
    from docx import Document
    from docx.shared import Inches, Pt

    document = Document()
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.space_after = Pt(0)
    document.styles["Normal"].paragraph_format.line_spacing = 1
    recognized = 0
    languages = set()
    previews = []
    empty_pages = []
    for page_number, lines, language, image in ocr_pdf_pages(source, workdir, progress, cancel):
        languages.add(language)
        if page_number > 1:
            document.add_page_break()
        document.add_heading(f"Page {page_number} - recognized text", level=1)
        useful = [line.strip() for line in lines if line.strip()]
        if useful:
            for line in useful:
                document.add_paragraph(line)
            recognized += 1
        else:
            empty_pages.append(page_number)
            document.add_paragraph("No text recognized on this page; see the original page below.")
        previews.append(image)
    if not recognized:
        raise ConversionError(
            "Local OCR could not recognize any text in this PDF. No DOCX was created."
        )
    document.add_page_break()
    document.add_heading("Original pages for reference", level=1)
    document.add_paragraph(
        "The preceding text was recognized automatically. Check it against these original pages; "
        "diagrams and complex layout have not been reconstructed as editable objects."
    )
    for index, image in enumerate(previews):
        check_cancel(cancel)
        if index:
            document.add_page_break()
        document.add_heading(f"Original page {index + 1}", level=2)
        # Keep every full page, including diagrams, visible on a reference page.
        with pymupdf.open(image) as raster:
            ratio = raster[0].rect.width / raster[0].rect.height
        document.add_picture(str(image), width=Inches(min(6, 7 * ratio)))
    check_cancel(cancel)
    progress(0.8, "Saving recognized text and original-page references")
    output = workdir / "output.docx"
    document.save(output)
    warnings = [
        (
            f"Local {backend_name()} OCR ({', '.join(sorted(languages))}) was used. Review recognized text for mistakes. "
            "Layout is simplified; original page images are included for reference."
        )
    ]
    if empty_pages:
        warnings.append("No text recognized on page(s): " + ", ".join(map(str, empty_pages)) + ".")
    return (EngineOutput(output, source.with_suffix(".docx").name, tuple(warnings)),)
